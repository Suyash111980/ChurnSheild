"""Bounded CSV/XLSX parsing and descriptive/predictive compatibility checks."""

from __future__ import annotations

import csv
import io
import math
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import numpy as np

from app.integrations.feature_contract import feature_columns, numeric_feature_columns

FEATURE_COLUMNS = feature_columns()
NUMERIC_FEATURES = numeric_feature_columns()

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_XLSX_FILES = 500
MAX_XLSX_COMPRESSION_RATIO = 200
REQUIRED_KEYS = ("id_student", "code_module", "code_presentation")
MAX_ISSUES = 100


class DatasetUploadError(ValueError):
    def __init__(self, message: str, *, status_code: int = 422):
        self.status_code = status_code
        super().__init__(message)


@dataclass
class ParsedDataset:
    frame: pd.DataFrame
    report: dict[str, Any]


def _read_csv(content: bytes) -> pd.DataFrame:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DatasetUploadError("CSV must use UTF-8 text encoding") from exc
    if "\x00" in text:
        raise DatasetUploadError("CSV contains invalid null bytes")
    try:
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            # Sniffer needs multiple distinguishable fields. A valid one-column
            # file should still receive the normal required-column feedback.
            dialect = csv.excel
        if dialect.delimiter != ",":
            raise DatasetUploadError("CSV must use comma-separated columns")
        header = next(csv.reader(io.StringIO(text), delimiter=",", strict=True), None)
        if header is None:
            raise DatasetUploadError("CSV must contain a header row")
        frame = pd.read_csv(io.StringIO(text), dtype="string", nrows=MAX_ROWS + 1)
        # pandas renames duplicate headers (for example, `id_student.1`). Keep
        # the submitted names so the validator can reject duplicates instead
        # of silently treating them as unexpected columns.
        if len(header) == len(frame.columns):
            frame.columns = header
    except DatasetUploadError:
        raise
    except Exception as exc:
        raise DatasetUploadError("CSV could not be parsed; check its header and quoting") from exc
    return frame


def _read_xlsx(content: bytes) -> pd.DataFrame:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_FILES:
                raise DatasetUploadError("Workbook contains too many internal files")
            uncompressed = sum(item.file_size for item in infos)
            if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise DatasetUploadError("Workbook expands beyond the allowed size")
            if any(item.file_size > 0 and item.compress_size == 0 for item in infos):
                raise DatasetUploadError("Workbook contains an invalid compressed entry")
            if any(
                item.compress_size and item.file_size / item.compress_size > MAX_XLSX_COMPRESSION_RATIO
                for item in infos
            ):
                raise DatasetUploadError("Workbook compression ratio exceeds the allowed limit")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise DatasetUploadError("File content is not a supported XLSX workbook")
            if any("vbaProject" in name for name in names):
                raise DatasetUploadError("Macro-enabled workbooks are not supported")
    except DatasetUploadError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise DatasetUploadError("XLSX workbook is invalid or corrupted") from exc

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        if not workbook.worksheets:
            raise DatasetUploadError("Workbook has no worksheets")
        sheet = workbook.worksheets[0]
        rows = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1, values_only=True),
            start=1,
        ):
            values = list(row)
            while values and values[-1] is None:
                values.pop()
            if len(values) > MAX_COLUMNS:
                raise DatasetUploadError(f"Workbook exceeds the {MAX_COLUMNS}-column limit")
            if row_number == 1:
                headers = [str(value).strip() if value is not None else "" for value in values]
                if not any(headers):
                    raise DatasetUploadError("Workbook header row is empty")
                continue
            if not any(value is not None for value in values):
                continue
            rows.append(values)
            if len(rows) > MAX_ROWS:
                raise DatasetUploadError(f"Workbook exceeds the {MAX_ROWS}-row limit")
        workbook.close()
        width = len(headers)
        if width > MAX_COLUMNS:
            raise DatasetUploadError(f"Workbook exceeds the {MAX_COLUMNS}-column limit")
        normalized = []
        for row in rows:
            if len(row) > width and any(value is not None for value in row[width:]):
                raise DatasetUploadError("Workbook contains values beyond its header columns")
            normalized.append((row + [None] * width)[:width])
        return pd.DataFrame(normalized, columns=headers, dtype="string")
    except DatasetUploadError:
        raise
    except Exception as exc:
        raise DatasetUploadError("XLSX workbook could not be read safely") from exc


def _json_value(value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def parse_dataset(filename: str, content: bytes) -> ParsedDataset:
    """Parse bounded file content and produce actionable validation metadata."""
    safe_name = Path(filename or "").name
    extension = Path(safe_name).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise DatasetUploadError("Only CSV and XLSX uploads are supported", status_code=415)
    if not content:
        raise DatasetUploadError("Uploaded file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise DatasetUploadError("Upload exceeds the 25 MiB file-size limit", status_code=413)
    is_zip = content.startswith(b"PK\x03\x04")
    if extension == ".xlsx" and not is_zip:
        raise DatasetUploadError("File extension and content do not match XLSX format")
    if extension == ".csv" and is_zip:
        raise DatasetUploadError("File extension and content do not match CSV format")
    frame = _read_xlsx(content) if extension == ".xlsx" else _read_csv(content)
    frame.columns = [str(column).strip() for column in frame.columns]
    issues: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    if not len(frame.columns) or any(not column for column in frame.columns):
        issues.append({"code": "empty_column_name", "message": "Column names must not be blank"})
    if len(frame.columns) != len(set(frame.columns)):
        issues.append({"code": "duplicate_columns", "message": "Column names must be unique"})
    if any(not column for column in frame.columns) or len(frame.columns) != len(set(frame.columns)):
        return ParsedDataset(
            frame=frame,
            report={
                "row_count": int(len(frame)),
                "column_count": int(len(frame.columns)),
                "columns": list(frame.columns),
                "required_columns": list(REQUIRED_KEYS),
                "missing_prediction_columns": list(FEATURE_COLUMNS),
                "descriptive_compatible": False,
                "prediction_compatible": False,
                "errors": issues[:MAX_ISSUES],
                "warnings": [],
                "preview": [],
            },
        )
    if len(frame.columns) > MAX_COLUMNS:
        raise DatasetUploadError(f"File exceeds the {MAX_COLUMNS}-column limit")
    if len(frame) > MAX_ROWS:
        raise DatasetUploadError(f"File exceeds the {MAX_ROWS}-row limit")
    if frame.empty:
        issues.append({"code": "no_rows", "message": "File must contain at least one data row"})
    for column in REQUIRED_KEYS:
        if column not in frame.columns:
            issues.append({"code": "missing_required_column", "column": column,
                           "message": f"Required column '{column}' is missing"})
    if all(column in frame for column in REQUIRED_KEYS):
        for column in REQUIRED_KEYS:
            normalized = frame[column].astype("string").str.strip()
            empty_count = int((normalized.isna() | normalized.eq("")).sum())
            if empty_count:
                issues.append({"code": "missing_key_values", "column": column,
                               "count": empty_count,
                               "message": f"{empty_count} rows have an empty {column}"})
            frame[column] = normalized
        duplicate_count = int(frame.duplicated(list(REQUIRED_KEYS), keep=False).sum())
        if duplicate_count:
            issues.append({"code": "duplicate_learner_course_keys", "count": duplicate_count,
                           "message": "Rows must be unique by learner, module, and presentation"})
    for column in NUMERIC_FEATURES:
        if column not in frame.columns:
            continue
        nonempty = frame[column].notna() & frame[column].astype("string").str.strip().ne("")
        converted = pd.to_numeric(frame[column], errors="coerce")
        invalid_mask = nonempty & (converted.isna() | ~np.isfinite(converted))
        invalid_count = int(invalid_mask.sum())
        if invalid_count:
            issues.append({"code": "invalid_numeric_values", "column": column,
                           "count": invalid_count,
                           "message": f"{invalid_count} values in {column} are not numeric"})
        if column in {"num_of_prev_attempts", "studied_credits"}:
            invalid_integer = converted.notna() & ((converted < 0) | (converted % 1 != 0))
            if invalid_integer.any():
                issues.append({"code": "invalid_integer_values", "column": column,
                               "count": int(invalid_integer.sum()),
                               "message": f"{column} values must be non-negative integers"})
        frame[column] = converted
    if "disability" in frame.columns:
        disability_values = frame["disability"].astype("string").str.strip().str.upper()
        invalid_disability = disability_values.notna() & ~disability_values.isin(["Y", "N", "TRUE", "FALSE", "1", "0"])
        if invalid_disability.any():
            issues.append({"code": "invalid_disability_values", "column": "disability",
                           "count": int(invalid_disability.sum()),
                           "message": "disability values must be Y/N, true/false, or 1/0"})
    if "target_withdrawn" in frame.columns:
        labels = frame["target_withdrawn"].astype("string").str.strip()
        invalid_labels = int((labels.notna() & ~labels.isin(["0", "1"])).sum())
        if invalid_labels:
            issues.append({"code": "invalid_target_values", "column": "target_withdrawn",
                           "count": invalid_labels,
                           "message": "target_withdrawn may contain only 0 or 1"})
    unexpected = sorted(set(frame.columns) - set(FEATURE_COLUMNS) - set(REQUIRED_KEYS) - {"target_withdrawn"}
                        - {"gender", "region", "highest_education", "imd_band", "age_band", "disability"})
    if unexpected:
        warnings.append({"code": "unexpected_columns", "columns": unexpected,
                         "message": "Unexpected columns are ignored during processing"})
    missing_features = [column for column in FEATURE_COLUMNS if column not in frame.columns]
    descriptive = all(column in frame.columns for column in REQUIRED_KEYS) and bool(len(frame))
    prediction = descriptive and not missing_features and not any(
        issue["code"] in {"invalid_numeric_values", "duplicate_columns", "duplicate_learner_course_keys"}
        for issue in issues
    )
    if missing_features:
        warnings.append({"code": "prediction_features_missing", "columns": missing_features,
                         "message": "Dataset is descriptive-compatible but lacks the complete model feature schema"})
    if "target_withdrawn" in frame.columns:
        warnings.append({"code": "target_column_ignored", "column": "target_withdrawn",
                         "message": "Outcome labels are retained for audit only and never passed to inference"})
    return ParsedDataset(
        frame=frame,
        report={
            "row_count": int(len(frame)),
            "column_count": int(len(frame.columns)),
            "columns": list(frame.columns),
            "required_columns": list(REQUIRED_KEYS),
            "missing_prediction_columns": missing_features,
            "descriptive_compatible": descriptive and not issues,
            "prediction_compatible": prediction and not issues,
            "errors": issues[:MAX_ISSUES],
            "warnings": warnings,
            "preview": [
                {column: _json_value(value) for column, value in row.items()}
                for row in frame.head(10).to_dict(orient="records")
            ],
        },
    )
